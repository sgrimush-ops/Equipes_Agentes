import pandas as pd
import numpy as np
from openpyxl.styles import PatternFill

print("Lendo a planilha sugestao.xlsx...")
df = pd.read_excel('sugestao.xlsx')

colunas_financeiras = ['Valor Total', 'Valor Desconto', 'Custo Liquido', 'Impostos']
for col in colunas_financeiras:
    if col in df.columns:
        df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)

niveis = ['DEPARTAMENTO', 'SECAO', 'GRUPO', 'SUBGRUPO', 'NIVEL5', 'NIVEL6']

for nivel in niveis:
    print(f"Calculando margem média ponderada para: {nivel}...")
    df[nivel] = df[nivel].replace(r'^\s*$', np.nan, regex=True)
    agrupado = df.groupby(nivel, dropna=True).agg({
        'Valor Total': 'sum',
        'Valor Desconto': 'sum',
        'Custo Liquido': 'sum',
        'Impostos': 'sum'
    }).reset_index()
    
    venda_real = agrupado['Valor Total'] - agrupado['Valor Desconto']
    lucro_liq = venda_real - agrupado['Custo Liquido'] - agrupado['Impostos']
    margem_perc = np.where(venda_real > 0, (lucro_liq / venda_real) * 100, 0)
    
    agrupado['MARGEM_CALCULADA'] = margem_perc
    mapa_margem = dict(zip(agrupado[nivel], agrupado['MARGEM_CALCULADA']))
    
    coluna_destino = f"{nivel}.1"
    # Adicionamos a coluna SEM o "if" de checagem. Se não existir, ela será criada.
    df[coluna_destino] = df[nivel].map(mapa_margem).round(2)

print("Aplicando regra de preenchimento para o NIVEL6...")
if 'NIVEL6.1' in df.columns and 'NIVEL5.1' in df.columns:
    df['NIVEL6.1'] = np.where((df['NIVEL6.1'].isna()) | (df['NIVEL6.1'] == 0), 
                              df['NIVEL5.1'], 
                              df['NIVEL6.1'])

print("Reorganizando colunas para garantir que DEPARTAMENTO esteja na posição correta...")
# Descobrir onde termina a base de dados (logo após a Lucrat.% Margem)
cols = list(df.columns)
if 'Lucrat.% Margem' in cols:
    idx_margem = cols.index('Lucrat.% Margem')
else:
    idx_margem = len(cols) - 7 # Tenta achar uma aproximação

cols_base = cols[:idx_margem+1]
cols_calculadas = ['DEPARTAMENTO.1', 'SECAO.1', 'GRUPO.1', 'SUBGRUPO.1', 'NIVEL5.1', 'NIVEL6.1']
cols_final = cols_base + [c for c in cols_calculadas if c in df.columns]
df = df[cols_final]

print("Limpando nomes de colunas...")
nova_lista_colunas = [c[:-2] if c.endswith('.1') else c for c in df.columns]
df.columns = nova_lista_colunas

nome_saida = 'sugestao_processada.xlsx'
print("Exportando e aplicando Filtros, Cores e Congelamento de Painéis...")

# Usando o openpyxl para manipular o visual do arquivo
writer = pd.ExcelWriter(nome_saida, engine='openpyxl')
df.to_excel(writer, index=False, sheet_name='Plan1')

worksheet = writer.sheets['Plan1']

# Congelar linha superior
worksheet.freeze_panes = "A2"

# Adicionar Autofiltro
worksheet.auto_filter.ref = worksheet.dimensions

# Definindo as cores
# Usaremos tons pasteis como no print do usuário
fill_laranja = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
fill_amarelo = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

for col_idx, col_name in enumerate(df.columns, 1): # 1-indexed for openpyxl
    fill = None
    if col_name in ['GRUPO', 'SUBGRUPO']:
        fill = fill_laranja
    elif col_name == 'NIVEL5':
        fill = fill_amarelo
        
    if fill:
        for row in range(2, len(df) + 2):
            worksheet.cell(row=row, column=col_idx).fill = fill

writer.close()
print(f"Sucesso! Planilha '{nome_saida}' formatada perfeitamente.")
