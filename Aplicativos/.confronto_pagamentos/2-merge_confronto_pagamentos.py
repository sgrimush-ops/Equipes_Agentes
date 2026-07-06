import os
import re
import pdfplumber
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def extrair_numero_base(val):
    if pd.isna(val) or not val:
        return ""
    val = str(val).strip()
    m = re.match(r'^(\d+)', val)
    return m.group(1) if m else val

def parse_br_float(val):
    if pd.isna(val) or not val:
        return 0.0
    val = str(val).strip()
    val = re.sub(r'[^\d,\.-]', '', val)
    if ',' in val and '.' in val:
        val = val.replace('.', '').replace(',', '.')
    elif ',' in val:
        val = val.replace(',', '.')
    try:
        return float(val)
    except:
        return 0.0

def processar_confronto():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    txt_path = os.path.join(current_dir, 'a_pagar.txt')
    pdf_path = os.path.join(current_dir, 'pagamentos.pdf')
    out_excel = os.path.join(current_dir, 'Confronto_Pagamentos_a_Pagar.xlsx')

    print("1. Carregando dados de a_pagar.txt...")
    df_txt = None
    for enc in ['utf-8', 'latin1', 'cp1252', 'utf-16']:
        try:
            df_txt = pd.read_csv(txt_path, sep=';', encoding=enc, dtype=str)
            break
        except Exception:
            continue

    if df_txt is None or df_txt.empty:
        raise ValueError("Falha ao carregar a_pagar.txt")

    df_txt['NRO_BASE_TXT'] = df_txt['TITULO'].apply(extrair_numero_base)
    df_txt['VALOR_TXT_NUM'] = df_txt['VALOR_PROJETADO'].apply(parse_br_float)

    print(f"-> {len(df_txt)} registros lidos do TXT.")

    print("2. Extraindo titulos do arquivo pagamentos.pdf...")
    pdf_records = []
    current_fornecedor = ""
    current_cod_forn = ""

    with pdfplumber.open(pdf_path) as pdf:
        for page_idx, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            for line in text.split('\n'):
                line = line.strip()
                if not line:
                    continue

                # Atualiza fornecedor
                if line.startswith("FORNECEDOR:") or line.startswith("Pessoa "):
                    clean_forn = re.sub(r'^(FORNECEDOR:|Pessoa)\s*', '', line).strip()
                    m_cod = re.match(r'^(\d+)\s+(.*)', clean_forn)
                    if m_cod:
                        current_cod_forn = m_cod.group(1)
                        current_fornecedor = m_cod.group(2)
                    else:
                        current_cod_forn = ""
                        current_fornecedor = clean_forn
                    continue

                parts = line.split()
                if len(parts) >= 6 and parts[0].isdigit():
                    dates = [p for p in parts if re.match(r'^\d{2}/\d{2}/\d{4}$', p)]
                    if len(dates) >= 1:
                        emp = parts[0]
                        especie = parts[1] if len(parts) > 1 else ""
                        titulo = parts[2] if len(parts) > 2 else ""
                        nro_doc = parts[3] if len(parts) > 3 else ""
                        
                        if re.match(r'^\d{2}/\d{2}/\d{4}$', nro_doc):
                            nro_doc = extrair_numero_base(titulo)
                        else:
                            nro_doc = extrair_numero_base(nro_doc)

                        dta_emissao = dates[0] if len(dates) >= 1 else ""
                        dta_vencimento = dates[1] if len(dates) >= 2 else ""

                        monetary_vals = [p for p in parts if ',' in p and re.match(r'^[\d\.,-]+$', p)]
                        vlr_nominal = parse_br_float(monetary_vals[-3]) if len(monetary_vals) >= 3 else 0.0
                        vlr_liquido = parse_br_float(monetary_vals[-1]) if len(monetary_vals) >= 1 else 0.0

                        pdf_records.append({
                            'EMP_PDF': emp,
                            'COD_FORN_PDF': current_cod_forn,
                            'FORNECEDOR_PDF': current_fornecedor,
                            'ESPECIE_PDF': especie,
                            'TITULO_PDF': titulo,
                            'NRODOC_BASE_PDF': nro_doc,
                            'DATA_EMISSAO_PDF': dta_emissao,
                            'DATA_VENCIMENTO_PDF': dta_vencimento,
                            'VALOR_NOMINAL_PDF': vlr_nominal,
                            'VALOR_LIQUIDO_PDF': vlr_liquido,
                            'LINHA_ORIGINAL': line
                        })

    df_pdf = pd.DataFrame(pdf_records)
    print(f"-> {len(df_pdf)} titulos extraidos do PDF.")

    print("3. Realizando o merge e confronto entre PDF e TXT...")
    bases_txt = set(df_txt['NRO_BASE_TXT'].dropna().unique()) - {"", "0"}
    
    df_pdf['ESTA_NO_TXT'] = df_pdf['NRODOC_BASE_PDF'].isin(bases_txt)
    df_so_pdf = df_pdf[~df_pdf['ESTA_NO_TXT']].copy()
    df_so_pdf = df_so_pdf.sort_values(by=['FORNECEDOR_PDF', 'DATA_VENCIMENTO_PDF'])

    df_ambos_pdf = df_pdf[df_pdf['ESTA_NO_TXT']].copy()
    
    df_txt_dedup = df_txt.drop_duplicates(subset=['NRO_BASE_TXT']).copy()
    df_conciliados = pd.merge(
        df_ambos_pdf, 
        df_txt_dedup[['NRO_BASE_TXT', 'TITULO', 'FORNECEDOR', 'VALOR_TXT_NUM', 'DATA_VENCIMENTO']], 
        left_on='NRODOC_BASE_PDF', 
        right_on='NRO_BASE_TXT', 
        how='left',
        suffixes=('_PDF', '_TXT')
    )

    bases_pdf = set(df_pdf['NRODOC_BASE_PDF'].dropna().unique()) - {"", "0"}
    df_txt['ESTA_NO_PDF'] = df_txt['NRO_BASE_TXT'].isin(bases_pdf)
    df_so_txt = df_txt[~df_txt['ESTA_NO_PDF']].copy()
    df_so_txt = df_so_txt.sort_values(by=['FORNECEDOR', 'DATA_VENCIMENTO'])

    print(f"   - Duplicatas SO NO PDF (Faltam no a_pagar.txt): {len(df_so_pdf)}")
    print(f"   - Duplicatas CONCILIADAS (Em ambos): {len(df_conciliados)}")
    print(f"   - Duplicatas SO NO TXT: {len(df_so_txt)}")

    print("4. Gerando relatorio Excel formatado...")
    wb = Workbook()
    
    font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    fill_header_red = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
    fill_header_green = PatternFill(start_color='375623', end_color='375623', fill_type='solid')
    fill_header_blue = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    font_bold = Font(name='Segoe UI', size=10, bold=True)
    font_regular = Font(name='Segoe UI', size=10)
    
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    border_total = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000')
    )

    # ABA 1: SÓ NO PDF (O principal pedido do usuário)
    ws1 = wb.active
    ws1.title = "Faltam no A_Pagar (So no PDF)"
    ws1.views.sheetView[0].showGridLines = True
    
    ws1.cell(row=1, column=1, value="DUPLICATAS EM ABERTO NO PDF QUE NAO ESTAO NO A_PAGAR.TXT").font = Font(name='Segoe UI', size=13, bold=True, color='C00000')
    ws1.cell(row=2, column=1, value="Títulos presentes na relação de pagamentos do banco/financeiro (pagamentos.pdf), mas não encontrados no relatório de compras (a_pagar.txt).").font = Font(name='Segoe UI', size=9, italic=True, color='595959')

    headers1 = [
        "Empresa", "Cód. Fornecedor", "Razão Social / Fornecedor (PDF)", 
        "Espécie", "Título PDF", "Nro. Doc Base", 
        "Data Emissão", "Data Vencimento", "Valor Nominal (R$)", "Valor Líquido (R$)"
    ]

    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_red
        cell.alignment = Alignment(horizontal='center' if col_idx in [1,2,4,6,7,8] else 'left', vertical='center')

    row_idx = 5
    for _, row in df_so_pdf.iterrows():
        ws1.cell(row=row_idx, column=1, value=row['EMP_PDF']).alignment = Alignment(horizontal='center')
        ws1.cell(row=row_idx, column=2, value=row['COD_FORN_PDF']).alignment = Alignment(horizontal='center')
        ws1.cell(row=row_idx, column=3, value=row['FORNECEDOR_PDF'])
        ws1.cell(row=row_idx, column=4, value=row['ESPECIE_PDF']).alignment = Alignment(horizontal='center')
        ws1.cell(row=row_idx, column=5, value=row['TITULO_PDF']).alignment = Alignment(horizontal='center')
        ws1.cell(row=row_idx, column=6, value=row['NRODOC_BASE_PDF']).alignment = Alignment(horizontal='center')
        ws1.cell(row=row_idx, column=7, value=row['DATA_EMISSAO_PDF']).alignment = Alignment(horizontal='center')
        ws1.cell(row=row_idx, column=8, value=row['DATA_VENCIMENTO_PDF']).alignment = Alignment(horizontal='center')
        
        c_nom = ws1.cell(row=row_idx, column=9, value=row['VALOR_NOMINAL_PDF'])
        c_nom.number_format = 'R$ #,##0.00'
        c_nom.alignment = Alignment(horizontal='right')
        
        c_liq = ws1.cell(row=row_idx, column=10, value=row['VALOR_LIQUIDO_PDF'])
        c_liq.number_format = 'R$ #,##0.00'
        c_liq.alignment = Alignment(horizontal='right')

        for c in range(1, 11):
            ws1.cell(row=row_idx, column=c).font = font_regular
            ws1.cell(row=row_idx, column=c).border = border_thin
        row_idx += 1

    ws1.cell(row=row_idx, column=8, value="TOTAL GERAL:").font = font_bold
    ws1.cell(row=row_idx, column=8).alignment = Alignment(horizontal='right')
    
    c_tot_nom = ws1.cell(row=row_idx, column=9, value=f"=SUM(I5:I{row_idx-1})")
    c_tot_nom.font = font_bold
    c_tot_nom.number_format = 'R$ #,##0.00'
    c_tot_nom.border = border_total

    c_tot_liq = ws1.cell(row=row_idx, column=10, value=f"=SUM(J5:J{row_idx-1})")
    c_tot_liq.font = font_bold
    c_tot_liq.number_format = 'R$ #,##0.00'
    c_tot_liq.border = border_total

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)
    ws1.column_dimensions['C'].width = 42

    # ABA 2: CONCILIADOS (Em Ambos)
    ws2 = wb.create_sheet(title="Conciliados (Em Ambos)")
    ws2.views.sheetView[0].showGridLines = True
    ws2.cell(row=1, column=1, value="DUPLICATAS CONCILIADAS (ENCONTRADAS TANTO NO PDF QUANTO NO A_PAGAR.TXT)").font = Font(name='Segoe UI', size=13, bold=True, color='375623')
    
    headers2 = [
        "Empresa", "Cód. Fornecedor", "Fornecedor (PDF)", "Espécie", 
        "Título PDF", "Título TXT", "Nro. Doc Base", 
        "Vencimento PDF", "Vencimento TXT", "Valor Líquido PDF (R$)", "Valor TXT (R$)"
    ]
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_green
        cell.alignment = Alignment(horizontal='center' if col_idx in [1,2,4,5,6,7,8,9] else 'left', vertical='center')

    row_idx = 4
    for _, row in df_conciliados.iterrows():
        ws2.cell(row=row_idx, column=1, value=row['EMP_PDF']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_idx, column=2, value=row['COD_FORN_PDF']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_idx, column=3, value=row['FORNECEDOR_PDF'])
        ws2.cell(row=row_idx, column=4, value=row['ESPECIE_PDF']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_idx, column=5, value=row['TITULO_PDF']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_idx, column=6, value=row['TITULO']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_idx, column=7, value=row['NRODOC_BASE_PDF']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_idx, column=8, value=row['DATA_VENCIMENTO_PDF']).alignment = Alignment(horizontal='center')
        ws2.cell(row=row_idx, column=9, value=row['DATA_VENCIMENTO']).alignment = Alignment(horizontal='center')
        
        c_p = ws2.cell(row=row_idx, column=10, value=row['VALOR_LIQUIDO_PDF'])
        c_p.number_format = 'R$ #,##0.00'
        
        c_t = ws2.cell(row=row_idx, column=11, value=row['VALOR_TXT_NUM'])
        c_t.number_format = 'R$ #,##0.00'

        for c in range(1, 12):
            ws2.cell(row=row_idx, column=c).font = font_regular
            ws2.cell(row=row_idx, column=c).border = border_thin
        row_idx += 1

    ws2.cell(row=row_idx, column=9, value="TOTAL CONCILIADO:").font = font_bold
    ws2.cell(row=row_idx, column=9).alignment = Alignment(horizontal='right')
    c_tot_p = ws2.cell(row=row_idx, column=10, value=f"=SUM(J4:J{row_idx-1})")
    c_tot_p.font = font_bold
    c_tot_p.number_format = 'R$ #,##0.00'
    c_tot_p.border = border_total
    
    c_tot_t = ws2.cell(row=row_idx, column=11, value=f"=SUM(K4:K{row_idx-1})")
    c_tot_t.font = font_bold
    c_tot_t.number_format = 'R$ #,##0.00'
    c_tot_t.border = border_total

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)
    ws2.column_dimensions['C'].width = 38

    # ABA 3: SÓ NO TXT (Faltam no PDF)
    ws3 = wb.create_sheet(title="Faltam no PDF (So no TXT)")
    ws3.views.sheetView[0].showGridLines = True
    ws3.cell(row=1, column=1, value="DUPLICATAS NO RELATORIO DO ERP (A_PAGAR.TXT) QUE NAO ESTAO NO PDF DO BANCO").font = Font(name='Segoe UI', size=13, bold=True, color='1F4E78')
    
    headers3 = [
        "Empresa", "Cód. Fornecedor", "Fornecedor (TXT)", "Título / NF", 
        "Nro. Doc Base", "Data Emissão", "Data Vencimento", "Status Pedido", "Valor (R$)"
    ]
    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_blue
        cell.alignment = Alignment(horizontal='center' if col_idx in [1,2,4,5,6,7,8] else 'left', vertical='center')

    row_idx = 4
    for _, row in df_so_txt.iterrows():
        ws3.cell(row=row_idx, column=1, value=row.get('EMPRESA', '')).alignment = Alignment(horizontal='center')
        ws3.cell(row=row_idx, column=2, value=row.get('COD_FORNECEDOR', '')).alignment = Alignment(horizontal='center')
        ws3.cell(row=row_idx, column=3, value=row.get('FORNECEDOR', ''))
        ws3.cell(row=row_idx, column=4, value=row.get('TITULO', '')).alignment = Alignment(horizontal='center')
        ws3.cell(row=row_idx, column=5, value=row.get('NRO_BASE_TXT', '')).alignment = Alignment(horizontal='center')
        ws3.cell(row=row_idx, column=6, value=row.get('DATA_EMISSAO', '')).alignment = Alignment(horizontal='center')
        ws3.cell(row=row_idx, column=7, value=row.get('DATA_VENCIMENTO', '')).alignment = Alignment(horizontal='center')
        ws3.cell(row=row_idx, column=8, value=row.get('STATUS_PEDIDO', '')).alignment = Alignment(horizontal='center')
        
        c_val = ws3.cell(row=row_idx, column=9, value=row['VALOR_TXT_NUM'])
        c_val.number_format = 'R$ #,##0.00'

        for c in range(1, 10):
            ws3.cell(row=row_idx, column=c).font = font_regular
            ws3.cell(row=row_idx, column=c).border = border_thin
        row_idx += 1

    ws3.cell(row=row_idx, column=8, value="TOTAL TXT:").font = font_bold
    ws3.cell(row=row_idx, column=8).alignment = Alignment(horizontal='right')
    c_tot_txt = ws3.cell(row=row_idx, column=9, value=f"=SUM(I4:I{row_idx-1})")
    c_tot_txt.font = font_bold
    c_tot_txt.number_format = 'R$ #,##0.00'
    c_tot_txt.border = border_total

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 12)
    ws3.column_dimensions['C'].width = 38

    wb.save(out_excel)
    print(f"\nRelatorio gerado com sucesso: {out_excel}")

if __name__ == "__main__":
    processar_confronto()
