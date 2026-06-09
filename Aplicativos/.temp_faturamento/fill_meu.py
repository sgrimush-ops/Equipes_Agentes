import os
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
import unicodedata

def clean_cnpj(val):
    if val is None:
        return ""
    val_str = str(val).strip()
    return re.sub(r'\D', '', val_str)

def clean_and_normalize(name):
    if not isinstance(name, str):
        return ""
    name = name.upper().strip()
    name = "".join(c for c in unicodedata.normalize('NFD', name) if unicodedata.category(c) != 'Mn')
    name = re.sub(r'[\.\,\/\-\_\(\)\&\+\:\#]', ' ', name)
    words = name.split()
    cleaned_words = []
    for w in words:
        if w in ['LTDA', 'SA', 'EIRELI', 'IND', 'COM', 'INDUSTRIA', 'COMERCIO', 'S', 'A', 'LTD', 'ME', 'EPP', 'FILIAL', 'CDD',
                 'DE', 'DA', 'DO', 'DOS', 'DAS', 'E', 'EM', 'PARA', 'POR']:
            continue
        if w == 'COOP':
            w = 'COOPERATIVA'
        elif w in ['LATICINIOS', 'LACTICINIOS', 'LACTICINIO', 'LATICINIO', 'LACTEA', 'LATSUL']:
            w = 'LATICINIO'
        elif w == 'CRBS':
            w = 'AMBEV'
        elif w == 'FRIBOI':
            w = 'JBS'
        elif w in ['DISTRIB', 'DIST']:
            w = 'DISTRIBUIDORA'
        elif w == 'COML':
            w = 'COMERCIAL'
        elif w == 'SUINOCULTORES':
            w = 'SUINOC'
        cleaned_words.append(w)
    
    res = " ".join(cleaned_words)
    if 'AMBEV' in res:
        return 'AMBEV'
    if 'JBS' in res:
        return 'JBS'
    return res

def main():
    dir_path = r"c:\Users\usr\Downloads\Equipes_Agentes\Aplicativos\.temp_faturamento"
    meu_path = os.path.join(dir_path, "meu.xlsx")
    fat_path = os.path.join(dir_path, "faturado2026.xlsx")
    
    if not os.path.exists(meu_path):
        print(f"Erro: Arquivo meu.xlsx não encontrado em {meu_path}")
        return
    if not os.path.exists(fat_path):
        print(f"Erro: Arquivo faturado2026.xlsx não encontrado em {fat_path}")
        return

    print(f"Carregando faturamento de {fat_path}...")
    df_fat = pd.read_excel(fat_path)
    
    col_cnpj_fat = 'CNPJ'
    col_desc_fat = 'DESCRICAO_FORNECEDOR'
    col_valor_fat = 'VALOR_TOTAL_FATURADO'
    
    # 1. Mapeamento de CNPJ -> Valor
    cnpj_map = {}
    for idx, row in df_fat.iterrows():
        if 'TOTAL GERAL' in str(row.iloc[0]).upper():
            continue
        cnpj_clean = clean_cnpj(row[col_cnpj_fat])
        val = row[col_valor_fat]
        try:
            val_float = float(val)
        except:
            val_float = 0.0
        if cnpj_clean:
            cnpj_map[cnpj_clean] = max(cnpj_map.get(cnpj_clean, 0.0), val_float)
            
    # 2. Mapeamento de Descrição Normalizada -> Soma dos Valores das filiais/registros
    name_to_rows = {}
    for idx, row in df_fat.iterrows():
        if 'TOTAL GERAL' in str(row.iloc[0]).upper():
            continue
        norm = clean_and_normalize(row[col_desc_fat])
        val = row[col_valor_fat]
        try:
            val_float = float(val)
        except:
            val_float = 0.0
        if norm:
            if norm not in name_to_rows:
                name_to_rows[norm] = []
            name_to_rows[norm].append((idx, val_float))
            
    name_map = {}
    for norm, rows in name_to_rows.items():
        unique_rows = {}
        for idx, val in rows:
            unique_rows[idx] = val
        name_map[norm] = sum(unique_rows.values())

    print(f"CNPJs únicos mapeados: {len(cnpj_map)}")
    print(f"Nomes normalizados únicos mapeados: {len(name_map)}")

    # Carregar meu.xlsx para atualizar
    print(f"Atualizando faturamento no arquivo {meu_path}...")
    wb = openpyxl.load_workbook(meu_path)
    ws = wb.active
    
    cnpj_col_idx = 1  # Coluna A (CNPJ)
    nome_col_idx = 2  # Coluna B (Nome Fantasia)
    fat_col_idx = 6   # Coluna F (Faturamento Últimos 12 Meses)
    
    font_data = Font(name="Calibri", size=11, bold=False)
    align_right = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    matched_cnpj = 0
    matched_name = 0
    not_matched_list = []
    valid_rows_count = 0
    total_rows = ws.max_row
    
    for r in range(2, total_rows + 1):
        cnpj_val = ws.cell(row=r, column=cnpj_col_idx).value
        nome_val = ws.cell(row=r, column=nome_col_idx).value
        
        # Se for linha vazia (CNPJ nulo/vazio), deixar a célula de faturamento nula
        if cnpj_val is None or str(cnpj_val).strip() == "":
            ws.cell(row=r, column=fat_col_idx, value=None)
            cell = ws.cell(row=r, column=fat_col_idx)
            cell.font = None
            cell.border = None
            cell.alignment = None
            cell.number_format = '@'
            continue
            
        cnpj_clean = clean_cnpj(cnpj_val)
        if not cnpj_clean:
            ws.cell(row=r, column=fat_col_idx, value=None)
            cell = ws.cell(row=r, column=fat_col_idx)
            cell.font = None
            cell.border = None
            cell.alignment = None
            cell.number_format = '@'
            continue
            
        valid_rows_count += 1
        fat_val = 0.0
        matched = False
        
        # 1. Tentar correspondência por CNPJ
        if cnpj_clean in cnpj_map:
            fat_val = cnpj_map[cnpj_clean]
            matched_cnpj += 1
            matched = True
            
        # 2. Se CNPJ falhar, tentar correspondência por Descrição do Nome Fantasia
        if not matched and nome_val:
            norm_meu = clean_and_normalize(nome_val)
            
            # Busca Exata
            if norm_meu in name_map:
                fat_val = name_map[norm_meu]
                matched_name += 1
                matched = True
            else:
                # Busca Parcial (Overlap de palavras)
                best_score = 0
                best_val = 0.0
                words_meu = set(norm_meu.split())
                
                if words_meu:
                    for k, v in name_map.items():
                        words_k = set(k.split())
                        if not words_k:
                            continue
                        intersect = words_meu.intersection(words_k)
                        if intersect:
                            score = len(intersect) / max(len(words_meu), len(words_k))
                            if words_meu.issubset(words_k) or words_k.issubset(words_meu):
                                score = max(score, 0.8)
                            if score > best_score:
                                best_score = score
                                best_val = v
                
                if best_score >= 0.5:
                    fat_val = best_val
                    matched_name += 1
                    matched = True
        
        if not matched:
            not_matched_list.append((nome_val, cnpj_val))

        cell = ws.cell(row=r, column=fat_col_idx, value=fat_val)
        cell.font = font_data
        cell.number_format = 'R$ #,##0.00'
        cell.alignment = align_right
        cell.border = thin_border
        
    # Habilitar linhas de grade
    ws.views.sheetView[0].showGridLines = True
    
    # Auto-ajustar largura das colunas
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value) if cell.value is not None else ""
            if isinstance(cell.value, float) or isinstance(cell.value, int):
                val_str = f"R$ {cell.value:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    wb.save(meu_path)
    wb.close()
    
    print(f"\n[SUCESSO] Planilha meu.xlsx atualizada com sucesso com Fallback por Descrição!")
    print(f"Linhas válidas identificadas: {valid_rows_count}")
    print(f"Correspondido por CNPJ: {matched_cnpj}")
    print(f"Correspondido por Nome Fantasia: {matched_name}")
    print(f"Total Correspondidos: {matched_cnpj + matched_name} de {valid_rows_count} ({(matched_cnpj + matched_name)/valid_rows_count*100:.2f}%)")
    
    if not_matched_list:
        print("\nNão correspondidos:")
        for nm, cnp in not_matched_list[:15]:
            print(f"  - '{nm}' | CNPJ: {cnp}")

if __name__ == "__main__":
    main()
