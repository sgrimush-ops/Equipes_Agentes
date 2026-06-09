import os
import sys
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import unicodedata
import re

def clean_and_normalize(name):
    if not isinstance(name, str):
        return ""
    name = name.upper().strip()
    # Remove accents
    name = "".join(c for c in unicodedata.normalize('NFD', name) if unicodedata.category(c) != 'Mn')
    # Replace punctuation and special characters with spaces
    name = re.sub(r'[\.\,\/\-\_\(\)\&\+\:\#]', ' ', name)
    words = name.split()
    cleaned_words = []
    for w in words:
        # Remove common business suffixes and prepositions
        if w in ['LTDA', 'SA', 'EIRELI', 'IND', 'COM', 'INDUSTRIA', 'COMERCIO', 'S', 'A', 'LTD', 'ME', 'EPP', 'FILIAL', 'CDD',
                 'DE', 'DA', 'DO', 'DOS', 'DAS', 'E', 'EM', 'PARA', 'POR']:
            continue
        if w == 'COOP':
            w = 'COOPERATIVA'
        elif w in ['LATICINIOS', 'LACTICINIOS', 'LACTICINIO', 'LATICINIO', 'LACTEA', 'LATSUL']:
            w = 'LATICINIO'
        elif w == 'CRBS':
            w = 'AMBEV'
        elif w == 'FRIBOI':
            w = 'JBS'
        elif w in ['DISTRIB', 'DIST']:
            w = 'DISTRIBUIDORA'
        elif w == 'COML':
            w = 'COMERCIAL'
        elif w == 'SUINOCULTORES':
            w = 'SUINOC'
        cleaned_words.append(w)
    
    res = " ".join(cleaned_words)
    # Consolidate major brands
    if 'AMBEV' in res:
        return 'AMBEV'
    if 'JBS' in res:
        return 'JBS'
    return res

def convert_txt_to_excel():
    print("=" * 60)
    print("  CONVERSOR DE FATURAMENTO: TXT -> EXCEL (COM COMPRA 2025)")
    print("=" * 60)

    # Definir diretórios de busca
    current_dir = Path(__file__).parent.resolve()
    
    # Procurar por qualquer arquivo .txt na pasta atual
    txt_files = list(current_dir.glob("*.txt"))
    if not txt_files:
        print("Erro: Nenhum arquivo .txt encontrado no diretório local.")
        sys.exit(1)
        
    # Usar o primeiro arquivo .txt encontrado (que será faturado2026.txt)
    txt_path = txt_files[0]
    print(f"Arquivo de origem encontrado: {txt_path.resolve()}")
    
    # Tentar ler com diferentes encodings
    df = None
    for encoding in ['utf-8-sig', 'latin-1', 'utf-8']:
        try:
            df = pd.read_csv(txt_path, sep=';', encoding=encoding, dtype=str)
            print(f"Leitura bem-sucedida usando encoding: {encoding}")
            break
        except Exception as e:
            continue
            
    if df is None or df.empty:
        print("Erro: Não foi possível ler o arquivo. Verifique o formato do arquivo TXT.")
        sys.exit(1)

    print(f"Dados carregados: {len(df)} linhas e {len(df.columns)} colunas.")
    
    # Padronizar nomes das colunas
    df.columns = [str(c).strip().upper() for c in df.columns]
    print(f"Colunas identificadas: {list(df.columns)}")
    
    col_valor = 'VALOR_TOTAL_FATURADO'
    col_codigo = 'CODIGO_FORNECEDOR'
    col_desc_2026 = 'DESCRICAO_FORNECEDOR'
    col_compra_2025 = 'COMPRA_2025'

    # Tratar os dados de 2026
    if col_valor in df.columns:
        df[col_valor] = df[col_valor].fillna('0')
        df[col_valor] = df[col_valor].astype(str).str.strip()
        df[col_valor] = df[col_valor].str.replace('R$', '', regex=False).str.strip()
        df[col_valor] = df[col_valor].str.replace('.', '', regex=False).str.replace(',', '.', regex=False)
        df[col_valor] = pd.to_numeric(df[col_valor], errors='coerce').fillna(0.0)
        print(f"Coluna de valor '{col_valor}' convertida para numérico.")
        
    if col_codigo in df.columns:
        try:
            df[col_codigo] = pd.to_numeric(df[col_codigo], errors='coerce').fillna(0).astype(int)
        except:
            pass
        print(f"Coluna de código '{col_codigo}' ajustada.")

    # Carregar faturado2025.xlsx e construir dicionário de busca
    f2025_path = current_dir / "faturado2025.xlsx"
    lookup = {}
    if f2025_path.exists():
        print(f"Carregando histórico de 2025: {f2025_path.resolve()}")
        try:
            df2025 = pd.read_excel(f2025_path)
            col_desc_2025_A = df2025.columns[0]
            col_desc_2025_B = df2025.columns[1]
            
            key_to_rows = {}
            for idx, row in df2025.iterrows():
                val = row['Compra']
                try:
                    if isinstance(val, str):
                        val = float(val.replace('.', '').replace(',', '.'))
                    else:
                        val = float(val)
                except:
                    val = 0.0
                    
                name_A = clean_and_normalize(row[col_desc_2025_A])
                name_B = clean_and_normalize(row[col_desc_2025_B])
                keys = set(filter(None, [name_A, name_B]))
                
                for k in keys:
                    if k not in key_to_rows:
                        key_to_rows[k] = []
                    key_to_rows[k].append((idx, val))
            
            # Consolidar soma de valores únicos por chave normalizada
            for k, rows in key_to_rows.items():
                unique_rows = {}
                for idx, val in rows:
                    unique_rows[idx] = val
                lookup[k] = sum(unique_rows.values())
            print(f"Histórico de 2025 carregado com {len(lookup)} chaves normalizadas.")
        except Exception as e:
            print(f"Erro ao ler faturado2025.xlsx: {e}")
    else:
        print(f"Aviso: {f2025_path.name} não encontrado no mesmo diretório. A coluna {col_compra_2025} será zerada.")

    # Buscar valores de Compra de 2025 para cada fornecedor de 2026
    compra_2025_values = []
    if lookup and col_desc_2026 in df.columns:
        print("Realizando cruzamento de fornecedores com dados de 2025...")
        matched_count = 0
        for idx, row in df.iterrows():
            raw_name = row[col_desc_2026]
            norm_name = clean_and_normalize(raw_name)
            
            # 1. Busca exata pela chave normalizada
            if norm_name in lookup:
                compra_2025_values.append(lookup[norm_name])
                matched_count += 1
            else:
                # 2. Busca parcial por overlap de palavras
                best_match_val = None
                best_score = 0
                norm_words = set(norm_name.split())
                
                if len(norm_words) >= 1:
                    for k, v in lookup.items():
                        k_words = set(k.split())
                        if not k_words:
                            continue
                        intersect = norm_words.intersection(k_words)
                        if intersect:
                            score = len(intersect) / max(len(norm_words), len(k_words))
                            if norm_words.issubset(k_words) or k_words.issubset(norm_words):
                                score = max(score, 0.8)
                            if score > best_score:
                                best_score = score
                                best_match_val = v
                
                if best_match_val is not None and best_score >= 0.5:
                    compra_2025_values.append(best_match_val)
                    matched_count += 1
                else:
                    compra_2025_values.append(0.0)
        print(f"Cruzamento concluído: {matched_count} de {len(df)} fornecedores correspondidos ({matched_count/len(df)*100:.2f}%).")
    else:
        compra_2025_values = [0.0] * len(df)
        
    df[col_compra_2025] = compra_2025_values

    # Gerar nome do Excel de saída
    excel_path = txt_path.with_suffix('.xlsx')
    
    # Salvar DataFrame temporário no Excel
    df.to_excel(excel_path, index=False, engine='openpyxl')
    
    # Formatação visual com openpyxl
    print("Aplicando formatação visual premium no Excel...")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    ws.title = "Faturamento"
    
    # Habilitar linhas de grade
    ws.views.sheetView[0].showGridLines = True
    
    # Estilos de design
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=11, bold=False)
    fill_header = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid") # Azul corporativo
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Formatar cabeçalho
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
        
    # Formatar linhas de dados
    for row_idx in range(2, len(df) + 2):
        # Zebrada: alternar fundo branco e cinza muito claro
        fill_data = PatternFill(start_color="F9FBFD" if row_idx % 2 == 0 else "FFFFFF", 
                                end_color="F9FBFD" if row_idx % 2 == 0 else "FFFFFF", 
                                fill_type="solid")
                                
        for col_idx in range(1, len(df.columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_data
            cell.fill = fill_data
            cell.border = thin_border
            
            # Alinhamentos e formatos específicos
            col_name = df.columns[col_idx - 1]
            if col_name in [col_valor, col_compra_2025]:
                cell.number_format = 'R$ #,##0.00'
                cell.alignment = align_right
            elif col_name == col_codigo or 'CNPJ' in col_name or 'CPF' in col_name:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Adicionar linha de total no final
    total_row_idx = len(df) + 2
    ws.cell(row=total_row_idx, column=1, value="TOTAL GERAL")
    ws.cell(row=total_row_idx, column=1).font = Font(name="Calibri", size=11, bold=True)
    ws.cell(row=total_row_idx, column=1).alignment = align_right
    
    # Preencher fundo da linha de total
    fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Azul clarinho
    
    # Colocar a fórmula da soma para colunas de valor
    for sum_col in [col_valor, col_compra_2025]:
        if sum_col in df.columns:
            val_col_idx = list(df.columns).index(sum_col) + 1
            col_letter = get_column_letter(val_col_idx)
            total_cell = ws.cell(row=total_row_idx, column=val_col_idx, value=f"=SUM({col_letter}2:{col_letter}{total_row_idx - 1})")
            total_cell.font = Font(name="Calibri", size=11, bold=True)
            total_cell.number_format = 'R$ #,##0.00'
            total_cell.alignment = align_right
    
    # Aplicar bordas e preenchimento na linha de total
    double_bottom_border = Border(
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='double', color='000000'),
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9')
    )
    
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=total_row_idx, column=col_idx)
        cell.fill = fill_total
        cell.border = double_bottom_border

    # Auto-ajustar largura das colunas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Se for fórmula, estimar tamanho
            val_str = str(cell.value)
            if val_str.startswith('='):
                val_str = "R$ 999.999.999,99" # Tamanho médio estimado do total
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Salvar alterações
    wb.save(excel_path)
    wb.close()
    
    print(f"\n[SUCESSO] Planilha Excel gerada e formatada com sucesso!")
    print(f"Local do arquivo: {excel_path.resolve()}")
    print("=" * 60)

if __name__ == "__main__":
    convert_txt_to_excel()