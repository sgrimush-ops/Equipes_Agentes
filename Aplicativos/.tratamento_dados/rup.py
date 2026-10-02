import os
import sys
from pathlib import Path

# Ajuste de encoding para o terminal Windows
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    try:
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def encontrar_query_parquet(base_dir: Path) -> Path:
    """Busca o arquivo query.parquet nos locais conhecidos do ecossistema."""
    candidatos = [
        base_dir.parent / "import_querys" / "query.parquet",
        base_dir.parent / "ProjetoBak_Sincronizador" / "bdados" / "query.parquet",
        base_dir.parent.parent / "import_querys" / "query.parquet",
        base_dir.parent / "bdados" / "query.parquet",
        base_dir.parent.parent / "bdados" / "query.parquet",
        base_dir / "query.parquet",
        base_dir.parent / "query.parquet",
    ]
    for c in candidatos:
        if c.exists():
            return c.resolve()
            
    # Busca recursiva no workspace
    raiz_workspace = base_dir.parent.parent
    if raiz_workspace.exists():
        encontrados = list(raiz_workspace.glob("**/query.parquet"))
        if encontrados:
            return max(encontrados, key=os.path.getmtime).resolve()
            
    raise FileNotFoundError("Arquivo query.parquet não encontrado nos diretórios do projeto.")

def formatar_aba_excel(ws, titulo_aba: str = ""):
    """Aplica estilização visual profissional na aba do Excel."""
    cor_cabecalho = "1B365D"  # Azul escuro corporativo
    fill_cabecalho = PatternFill(start_color=cor_cabecalho, end_color=cor_cabecalho, fill_type="solid")
    font_cabecalho = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    
    font_corpo = Font(name="Segoe UI", size=9)
    font_negrito = Font(name="Segoe UI", size=9, bold=True)
    
    fill_zebrado = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_cd_disponivel = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Verde claro para CD > 0
    font_cd_disponivel = Font(name="Segoe UI", size=9, color="276A3C", bold=True)
    
    border_fina = Side(style='thin', color="D3D3D3")
    borda_padrao = Border(left=border_fina, right=border_fina, top=border_fina, bottom=border_fina)
    
    # 1. Cabeçalho
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = fill_cabecalho
        cell.font = font_cabecalho
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = borda_padrao
    
    ws.row_dimensions[1].height = 26
    
    # Identificar índices de colunas relevantes
    headers = [str(ws.cell(row=1, column=col).value).upper() for col in range(1, ws.max_column + 1)]
    col_estoque_cd = None
    if "ESTOQUE CD" in headers:
        col_estoque_cd = headers.index("ESTOQUE CD") + 1
    elif "ESTOQUE_CD" in headers:
        col_estoque_cd = headers.index("ESTOQUE_CD") + 1
        
    # 2. Dados
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 20
        is_par = (row_idx % 2 == 0)
        
        # Verificar se CD tem estoque
        cd_tem_estoque = False
        if col_estoque_cd:
            val_cd = ws.cell(row=row_idx, column=col_estoque_cd).value
            try:
                if val_cd and float(val_cd) > 0:
                    cd_tem_estoque = True
            except (ValueError, TypeError):
                pass

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_corpo
            cell.border = borda_padrao
            
            # Zebrado básico
            if is_par:
                cell.fill = fill_zebrado
                
            header_name = headers[col_idx - 1] if col_idx - 1 < len(headers) else ""
            
            # Formatação por tipo de coluna
            if any(k in header_name for k in ["DATA CADASTRO", "DATA_CADASTRO", "DTA CADASTRO", "DTA_CADASTRO"]):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif any(k in header_name for k in ["SEQPRODUTO", "CODIGO", "LOJA", "EMPRESA", "DIAS"]):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
            elif any(k in header_name for k in ["DESC", "FORNECEDOR", "COMPRADOR", "STATUS", "FORMA"]):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif any(k in header_name for k in ["ESTOQUE DA LOJA", "ESTOQUE_LOJA"]):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00" if cell.value % 1 != 0 else "#,##0"
            elif any(k in header_name for k in ["ESTOQUE CD", "ESTOQUE_CD"]):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00" if cell.value % 1 != 0 else "#,##0"
                if cd_tem_estoque:
                    cell.fill = fill_cd_disponivel
                    cell.font = font_cd_disponivel
            elif any(k in header_name for k in ["VENDA DIARIA", "VENDA_DIARIA"]):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
            elif any(k in header_name for k in ["VENDA SEMANAL", "VENDA_SEMANAL", "VENDA 90 DIAS", "VENDA_90_DIAS", "POTENCIAL"]):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
                    
    # Congelar cabeçalho e ativar filtros
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    
    # Autoajuste de largura de colunas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row == 1:
                max_len = max(max_len, len(val_str) + 4)
            else:
                max_len = max(max_len, min(len(val_str) + 2, 60))
        ws.column_dimensions[col_letter].width = max(max_len, 12)

def criar_relatorio_ruptura():
    dir_script = Path(__file__).parent.resolve()
    parquet_path = encontrar_query_parquet(dir_script)
    saida_excel = dir_script / "ruptura_joao_batista_lojas_11_12_13.xlsx"
    
    print(f"[1/5] Carregando catálogo: {parquet_path}")
    df = pd.read_parquet(parquet_path)
    
    # Saneamento de colunas numéricas
    cols_numericas = ['QUANTIDADE_DISPONIVEL', 'QTD_VENDIDA_PERIODO', 'DIAS_PESQUISA', 'CODIGO_PRODUTO', 'CODIGO_EMPRESA']
    for col in ['QUANTIDADE_DISPONIVEL', 'QTD_VENDIDA_PERIODO']:
        if col in df.columns:
            df[col] = df[col].astype(str).str.replace(',', '.', regex=False)
            df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)
            
    if 'DIAS_PESQUISA' in df.columns:
        df['DIAS_PESQUISA'] = pd.to_numeric(df['DIAS_PESQUISA'], errors='coerce').fillna(90).astype(int)
    else:
        df['DIAS_PESQUISA'] = 90
        
    print("[2/5] Filtrando produtos do comprador JOAO BATISTA...")
    df_jb = df[df['COMPRADOR'].astype(str).str.upper().str.strip() == 'JOAO BATISTA'].copy()
    
    if df_jb.empty:
        print("[ERRO] Nenhum produto encontrado para o comprador JOAO BATISTA!")
        return
        
    # Mapeamento do Estoque CD (Empresa 15 é o CD de João Batista)
    print("[3/5] Mapeando estoque disponível no CD 15...")
    df_cd15 = df_jb[df_jb['CODIGO_EMPRESA'] == 15]
    cd_map = df_cd15.set_index('CODIGO_PRODUTO')['QUANTIDADE_DISPONIVEL'].to_dict()
    
    # Lojas solicitadas: 11, 12, 13
    lojas_alvo = [11, 12, 13]
    df_lojas = df_jb[df_jb['CODIGO_EMPRESA'].isin(lojas_alvo)].copy()
    
    # Filtro de ruptura da loja: QUANTIDADE_DISPONIVEL <= 0
    print("[4/5] Filtrando produtos em ruptura nas lojas 11, 12, 13...")
    df_rup = df_lojas[df_lojas['QUANTIDADE_DISPONIVEL'] <= 0].copy()
    
    # Cálculos das colunas solicitadas
    df_rup['ESTOQUE_CD'] = df_rup['CODIGO_PRODUTO'].map(cd_map).fillna(0)
    
    # Data de Cadastro do Item e Contagem de Dias do Cadastro até Hoje
    hoje = pd.Timestamp.now().normalize()
    if 'DATA_CADASTRO_PRODUTO' in df_rup.columns:
        dt_cad = pd.to_datetime(df_rup['DATA_CADASTRO_PRODUTO'], format='mixed', dayfirst=True, errors='coerce')
        df_rup['DATA_CADASTRO'] = dt_cad.dt.strftime('%d/%m/%Y').fillna(df_rup['DATA_CADASTRO_PRODUTO'].astype(str))
        df_rup['DIAS_CADASTRO'] = (hoje - dt_cad).dt.days.fillna(0).astype(int)
    else:
        df_rup['DATA_CADASTRO'] = '-'
        df_rup['DIAS_CADASTRO'] = 0
    
    dias_pesq = df_rup['DIAS_PESQUISA'].replace(0, 90)
    df_rup['VENDA_DIARIA'] = (df_rup['QTD_VENDIDA_PERIODO'] / dias_pesq).round(4)
    df_rup['VENDA_SEMANAL'] = (df_rup['VENDA_DIARIA'] * 7.0).round(2)
    df_rup['VENDA_90_DIAS'] = df_rup['QTD_VENDIDA_PERIODO'].round(2)
    
    # Ordenar registros por Loja e por maior venda semanal (prioridade de reposição)
    df_rup = df_rup.sort_values(by=['CODIGO_EMPRESA', 'VENDA_SEMANAL', 'CODIGO_PRODUTO'], ascending=[True, False, True])
    
    # Colunas canônicas formatadas
    colunas_exportacao = {
        'CODIGO_EMPRESA': 'LOJA',
        'CODIGO_PRODUTO': 'SEQPRODUTO',
        'DESCRICAO_PRODUTO': 'DESCCMPLETA',
        'DATA_CADASTRO': 'DATA CADASTRO',
        'DIAS_CADASTRO': 'DIAS CADASTRO',
        'QUANTIDADE_DISPONIVEL': 'ESTOQUE DA LOJA',
        'ESTOQUE_CD': 'ESTOQUE CD',
        'VENDA_DIARIA': 'VENDA DIARIA',
        'VENDA_SEMANAL': 'VENDA SEMANAL',
        'VENDA_90_DIAS': 'VENDA 90 DIAS',
        'STATUS_COMPRA': 'STATUS COMPRA',
        'FORMA_ABASTECIMENTO': 'FORMA ABASTECIMENTO',
        'FORNECEDOR': 'FORNECEDOR'
    }
    
    df_export_geral = df_rup[[c for c in colunas_exportacao.keys() if c in df_rup.columns]].rename(columns=colunas_exportacao)
    
    print(f"[5/5] Gerando arquivo Excel formatado: {saida_excel.name}...")
    with pd.ExcelWriter(saida_excel, engine='openpyxl') as writer:
        # 1. Aba Geral Consolidada
        df_export_geral.to_excel(writer, sheet_name="Ruptura Geral", index=False)
        
        # 2. Abas Individuais por Loja
        for lj in lojas_alvo:
            df_lj = df_rup[df_rup['CODIGO_EMPRESA'] == lj]
            df_lj_export = df_lj[[c for c in colunas_exportacao.keys() if c in df_lj.columns]].rename(columns=colunas_exportacao)
            # Remove coluna da loja na aba da própria loja
            cols_loja_aba = [c for c in df_lj_export.columns if c != 'LOJA']
            df_lj_export[cols_loja_aba].to_excel(writer, sheet_name=f"Loja {lj}", index=False)
            
        # 3. Aba Resumo Executivo
        resumo_data = []
        for lj in lojas_alvo:
            df_lj = df_rup[df_rup['CODIGO_EMPRESA'] == lj]
            qtd_rup = len(df_lj)
            qtd_com_cd = (df_lj['ESTOQUE_CD'] > 0).sum()
            qtd_sem_cd = (df_lj['ESTOQUE_CD'] <= 0).sum()
            venda_sem_perda = df_lj['VENDA_SEMANAL'].sum()
            resumo_data.append({
                'LOJA': f"Loja {lj}",
                'QTD PRODUTOS EM RUPTURA': qtd_rup,
                'RUPTURAS COM ESTOQUE NO CD (Abastecimento Imediato)': qtd_com_cd,
                'RUPTURAS SEM ESTOQUE NO CD': qtd_sem_cd,
                'POTENCIAL DE VENDA SEMANAL (Itens em Ruptura)': round(venda_sem_perda, 2)
            })
        
        df_resumo = pd.DataFrame(resumo_data)
        linha_total = pd.DataFrame([{
            'LOJA': 'TOTAL CONSOLIDADO',
            'QTD PRODUTOS EM RUPTURA': df_resumo['QTD PRODUTOS EM RUPTURA'].sum(),
            'RUPTURAS COM ESTOQUE NO CD (Abastecimento Imediato)': df_resumo['RUPTURAS COM ESTOQUE NO CD (Abastecimento Imediato)'].sum(),
            'RUPTURAS SEM ESTOQUE NO CD': df_resumo['RUPTURAS SEM ESTOQUE NO CD'].sum(),
            'POTENCIAL DE VENDA SEMANAL (Itens em Ruptura)': round(df_resumo['POTENCIAL DE VENDA SEMANAL (Itens em Ruptura)'].sum(), 2)
        }])
        df_resumo = pd.concat([df_resumo, linha_total], ignore_index=True)
        df_resumo.to_excel(writer, sheet_name="Resumo Indicadores", index=False)

    # Reabrir para aplicar estilos visuais com openpyxl
    wb = openpyxl.load_workbook(saida_excel)
    
    # Reordenar para que a aba de Resumo seja a primeira
    todas_abas = wb.sheetnames
    if "Resumo Indicadores" in todas_abas:
        idx_resumo = todas_abas.index("Resumo Indicadores")
        wb._sheets.insert(0, wb._sheets.pop(idx_resumo))
        
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        formatar_aba_excel(ws, sheet_name)
        
    wb.save(saida_excel)
    
    print("\n" + "="*70)
    print(f"✅ RELATÓRIO DE RUPTURA GERADO COM SUCESSO!")
    print(f"📁 Arquivo: {saida_excel}")
    print(f"📊 Total de ocorrências de ruptura: {len(df_rup)}")
    print(f"📦 Produtos distintos afetados: {df_rup['CODIGO_PRODUTO'].nunique()}")
    print("="*70 + "\n")
    
    for lj in lojas_alvo:
        sub = df_rup[df_rup['CODIGO_EMPRESA'] == lj]
        com_cd = (sub['ESTOQUE_CD'] > 0).sum()
        print(f"  • Loja {lj:02d}: {len(sub):4d} itens em ruptura | {com_cd:4d} com saldo disponível no CD15")
    print("="*70)

if __name__ == "__main__":
    criar_relatorio_ruptura()
